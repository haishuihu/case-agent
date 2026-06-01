from io import BytesIO

from openpyxl import load_workbook

from app.schemas.testcase import GenerationMetadata, TestCaseItem
from app.services.testcase_exporter import (
    _build_filename,
    _content_disposition,
    _sanitize_filename,
    build_markdown_table,
    build_skill_template_markdown,
    build_xlsx_bytes,
)


def _make_case(summary: str, priority: str = "P0", module: str = "登录模块", category: str = "正常流程") -> TestCaseItem:
    return TestCaseItem(summary=summary, priority=priority, module=module, category=category)


def test_build_xlsx_bytes_contains_cases_and_metadata() -> None:
    cases = [
        _make_case("验证登录成功：已注册用户 → 输入手机号和验证码，点击登录 → 登录成功并进入首页")
    ]
    metadata = GenerationMetadata(
        caseCount=1,
        categoryCounts={"正常流程": 1},
        moduleCounts={"登录模块": 1},
        priorityCounts={"P0": 1},
        warnings=["正常流程批次达到输出长度上限，结果可能不完整。"],
        truncated=True,
        incompleteCategories=["正常流程"],
    )

    workbook = load_workbook(BytesIO(build_xlsx_bytes(cases, metadata)))

    assert "测试用例" in workbook.sheetnames
    assert "生成信息" in workbook.sheetnames
    assert workbook["测试用例"]["A2"].value == "P0"
    assert workbook["测试用例"]["B2"].value == "登录模块"
    assert workbook["测试用例"]["C2"].value == "正常流程"
    assert workbook["生成信息"]["B2"].value == 1


def test_build_skill_template_follows_sections() -> None:
    cases = [
        _make_case("验证登录成功：用户已注册 → 输入正确手机号和验证码，点击登录 → 登录成功并进入首页", "P0"),
        _make_case("验证密码错误提示：用户已注册 → 输入错误密码，点击登录 → 提示「密码错误」", "P2", "登录模块", "异常与错误处理"),
    ]
    metadata = GenerationMetadata(
        caseCount=2,
        categoryCounts={"正常流程": 1, "异常与错误处理": 1},
        moduleCounts={"登录模块": 2},
        priorityCounts={"P0": 1, "P2": 1},
        reviewQuestions=[],
        warnings=[],
        truncated=False,
        incompleteCategories=[],
    )
    md = build_skill_template_markdown(
        cases,
        metadata,
        requirement_text="用户登录",
        document_titles=["rules.txt"],
        modules_overview=[("M01", "登录模块", "登录与鉴权")],
    )
    assert "# 【登录模块】功能测试用例" in md
    assert "> 编制人：" in md
    assert "> 日期：" in md
    assert "> 依据文档：rules.txt" in md
    assert "## 一、功能模块概述" in md
    assert "### 2.1 模块清单表格" in md
    assert "## 二、冒烟测试用例" in md
    assert "### 2.1 主流程概述" in md
    assert "### 2.2 冒烟测试用例" in md
    assert "### 2.3 冒烟测试快速检查清单" in md
    assert "### 2.4 冒烟测试统计" in md
    assert "## 三、测试用例明细" in md
    assert "## 四、异常场景测试用例" in md
    assert "## 九、测试用例统计" in md
    assert "## 十、待确认问题标记说明" in md
    assert "## 十一、评审记录" in md
    assert "### 产品经理评审" in md
    assert "## 十二、文档版本记录" in md
    assert "## 附录：格式规范" in md


def test_build_markdown_table_includes_priority_and_module_columns() -> None:
    markdown = build_markdown_table(
        [_make_case("验证登录成功：已注册用户 → 输入正确手机号和验证码，点击登录 → 登录成功并进入首页")]
    )

    assert "| 优先级 | 模块 | 分类 |" in markdown
    assert "P0" in markdown
    assert "登录模块" in markdown
    assert "正常流程" in markdown


def test_sanitize_filename_removes_illegal_chars() -> None:
    assert _sanitize_filename("商品营销：视频生成") == "商品营销视频生成"
    assert _sanitize_filename("test/file") == "testfile"
    assert _sanitize_filename("a:b*c?d<e>|f") == "abcdef"
    assert _sanitize_filename("normal name") == "normal name"


def test_build_filename_from_title() -> None:
    assert _build_filename("商品营销视频生成", "测试用例.md") == "商品营销视频生成测试用例.md"
    assert _build_filename("A" * 30, ".csv") == ("A" * 30) + ".csv"
    assert _build_filename("", ".xlsx") == "测试用例.xlsx"
    assert _build_filename(None, ".md") == "测试用例.md"
    assert _build_filename("大师创作资产改为资产库.md", "测试用例.md") == "大师创作资产改为资产库测试用例.md"
    assert _build_filename("需求文档.CSV", "测试用例.csv") == "需求文档测试用例.csv"
    assert _build_filename("测试.TXT", "测试用例.txt") == "测试测试用例.txt"


def test_content_disposition_is_ascii_safe() -> None:
    import re

    header = _content_disposition("商品营销测试用例.md")
    assert "filename*=UTF-8''" in header
    m = re.search(r'filename="([^"]+)"', header)
    assert m and m.group(1) == "%E5%95%86%E5%93%81%E8%90%A5%E9%94%80%E6%B5%8B%E8%AF%95%E7%94%A8%E4%BE%8B.md"
